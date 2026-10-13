"""
Teste de regressão: tempos exibidos em minutos e segundos (adendo 10.75).

Antes: "21.8" (minuto decimal) na tela/PDF dos relatórios, "21m 48s" no KPI
do dashboard, minutos com 2 casas no Excel do "Tempos por senha". Agora uma
função única (`app/utils/tempo.py::formatar_duracao`, filtro Jinja
`duracao`): "mm:ss" abaixo de 1 hora, "h:mm:ss" a partir de 1 hora, segundos
arredondados a partir dos segundos EXATOS; sem dado = vazio; negativo com
sinal. Excel: duração nativa (segundos arredondados / 86400, formato
[h]:mm:ss). CSV: minutos decimais + a mesma duração em texto.

Cenário (08/10/2026, Brasília):
  A - espera 21 min 48 s (1308 s) e atendimento 975,4 min (16:15:24)
  B - (outra fila/operador) atendimento NEGATIVO de 5 s (fim registrado antes
      da chamada - relógio fora de ordem): exibido "-00:05", sem esconder

Execução: python -m tests.test_regressao_duracao (a partir da raiz do
projeto, com o virtualenv ativado).
"""
import os

os.environ.setdefault("SECRET_KEY", "test")
os.environ.setdefault("FLASK_ENV", "testing")

import html
import io
import re
import zipfile
from datetime import datetime, timedelta
from unittest.mock import patch
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

from app import create_app
from app.extensions import db
from app.models import Empresa, Fila, LogAtendimento, Perfil, Senha, Unidade, Usuario
from app.relatorios.routes import _duracao_excel, _resumo_tempos
from app.utils.tempo import formatar_duracao, segundos_arredondados

UTC = ZoneInfo("UTC")
app = create_app("testing")

passos = []


def ok(msg):
    passos.append(msg)
    print("OK:", msg)


# ===========================================================================
# 1) A função única de formatação.
# ===========================================================================
casos = {
    21.8 * 60: "21:48",          # o caso relatado (0,8 min = 48 s)
    59 * 60 + 59.6: "1:00:00",   # arredondamento que cruza a hora
    975.4 * 60: "16:15:24",
    125: "02:05",
    21.5: "00:22",               # meio para cima
    20.5: "00:21",               # meio para cima (o round() do Python daria 20: empate vai para o par)
    0: "00:00",
    -5: "-00:05",
    None: "",
}
for segundos, esperado in casos.items():
    assert formatar_duracao(segundos) == esperado, (segundos, formatar_duracao(segundos), esperado)
assert segundos_arredondados(20.5) == 21 and segundos_arredondados(None) is None
assert app.jinja_env.filters["duracao"] is formatar_duracao
ok("formatar_duracao: 21:48, 1:00:00 (cruza a hora), 16:15:24, 02:05, meio para cima, 00:00, -00:05, None -> ''")

# Resumo sobre os segundos EXATOS (antes: média do minuto já arredondado a 2 casas).
linhas_resumo = [{f"{c}_segundos": s for c in ("espera_inicial", "espera_total", "atendimento_total")}
                 for s in (10, 11, 11)]
resumo = _resumo_tempos(linhas_resumo)["espera_inicial"]
assert abs(resumo["media"] - 32 / 3) < 1e-9 and resumo["mediana"] == 11 and resumo["quantidade"] == 3
media_antiga = sum(round(s / 60, 2) for s in (10, 11, 11)) / 3 * 60  # 10,6 s
assert abs(media_antiga - resumo["media"]) > 0.05  # a correção muda o valor (< 1 s)
ok("Resumo: média/mediana sobre os segundos exatos (10,667 s), não sobre minutos arredondados (10,6 s)")

# ===========================================================================
# Cenário.
# ===========================================================================
T0 = datetime(2026, 10, 8, 12, 0)  # 09:00 em Brasília
with app.app_context():
    db.create_all()
    perfil = Perfil(nome="admin")
    empresa = Empresa(nome="Clinica Duracao")
    db.session.add_all([perfil, empresa])
    db.session.commit()
    unidade = Unidade(empresa_id=empresa.id, nome="Matriz", sigla="MAT")
    db.session.add(unidade)
    db.session.commit()
    fila_a = Fila(empresa_id=empresa.id, unidade_id=unidade.id, nome="Coleta")
    fila_b = Fila(empresa_id=empresa.id, unidade_id=unidade.id, nome="Triagem")
    db.session.add_all([fila_a, fila_b])
    db.session.commit()
    admin = Usuario(empresa_id=empresa.id, nome="Admin", login="admin_duracao")
    admin.set_senha("senha123")
    admin.perfis = [perfil]
    op_b = Usuario(empresa_id=empresa.id, nome="Operador B", login="op_b_duracao")
    op_b.set_senha("senha123")
    db.session.add_all([admin, op_b])
    db.session.commit()

    def senha(codigo, fila, eventos, usuario):
        s = Senha(empresa_id=empresa.id, unidade_id=unidade.id, fila_id=fila.id, codigo=codigo,
                  sequencial=1, prioridade=False, status="finalizada", criada_em=T0)
        db.session.add(s)
        db.session.commit()
        for evento, segundos in eventos:
            db.session.add(LogAtendimento(
                senha_id=s.id, evento=evento, timestamp=T0 + timedelta(seconds=segundos),
                usuario_id=None if evento == "gerada" else usuario.id, unidade_id=unidade.id,
                fila_id=fila.id, setor="recepcao",
            ))
        db.session.commit()

    senha("A", fila_a, [("gerada", 0), ("chamada_senha", 1308), ("fim_atendimento", 1308 + 975.4 * 60)], admin)
    senha("B", fila_b, [("gerada", 0), ("chamada_senha", 60), ("fim_atendimento", 55)], op_b)
    ids = dict(empresa=empresa.id, unidade=unidade.id)

c = app.test_client()
assert c.post("/login", data={"login": "admin_duracao", "senha": "senha123"}, follow_redirects=True).status_code == 200
PERIODO = "data_inicio=2026-10-08&data_fim=2026-10-08"


def tela(rota):
    resp = c.get(f"/relatorios/{rota}?{PERIODO}")
    assert resp.status_code == 200, (rota, resp.status_code)
    return html.unescape(resp.get_data(as_text=True))


def baixar(rota, formato):
    resp = c.get(f"/relatorios/{rota}?{PERIODO}&formato={formato}")
    assert resp.status_code == 200, (rota, formato, resp.status_code)
    resp.direct_passthrough = False
    return resp.get_data()


def texto_pdf(conteudo):
    return " ".join("".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(conteudo)).pages).split())


# ===========================================================================
# 2) Tela e PDF: "Tempos por senha" e os relatórios agregados.
# ===========================================================================
t = tela("tempos-por-senha")
assert "<th>Espera inicial (mm:ss)</th>" in t and "<th>Atendimento total (mm:ss)</th>" in t
assert "<td>21:48</td>" in t and "<td>16:15:24</td>" in t
assert "<td>-00:05</td>" in t  # negativo exibido com sinal, sem esconder
assert "Média: 11:24" in t  # espera inicial média (1308 + 60) / 2 = 684 s
assert "No Excel, os tempos são durações (h:mm:ss)" in t
assert not re.search(r"<td>\d+\.\d</td>", t)  # nenhum minuto decimal na tabela
ok("Tempos por senha (tela): 21:48, 16:15:24, -00:05, média 11:24 nos cartões, cabeçalhos (mm:ss), nota do Excel")

for rota, esperado in (("tempo-espera", "21:48"), ("tempo-atendimento", "16:15:24"), ("ranking-operadores", "16:15:24")):
    pagina = tela(rota)
    assert "<th>Tempo médio (mm:ss)</th>" in pagina, rota
    assert f"<td>{esperado}</td>" in pagina, (rota, esperado)
    assert " min</td>" not in pagina, rota
assert "<td>-00:05</td>" in tela("tempo-atendimento")
ok("Relatórios agregados (tela): tempo de espera 21:48, tempo de atendimento e ranking 16:15:24, negativo -00:05")

pdf = texto_pdf(baixar("tempos-por-senha", "pdf"))
assert "21:48" in pdf and "16:15:24" in pdf and "11:24" in pdf
pdf_espera = texto_pdf(baixar("tempo-espera", "pdf"))
assert "21:48" in pdf_espera and "Tempo médio (mm:ss)" in pdf_espera
ok("PDF (Chromium real): Tempos por senha e Tempo de espera com mm:ss / h:mm:ss")

# ===========================================================================
# 3) Excel: duração NATIVA - célula numérica, valor = segundos/86400, [h]:mm:ss.
# ===========================================================================
conteudo = baixar("tempos-por-senha", "xlsx")
livro = openpyxl.load_workbook(io.BytesIO(conteudo))
xml_abas = {}
with zipfile.ZipFile(io.BytesIO(conteudo)) as z:
    for i, nome in enumerate(livro.sheetnames, start=1):
        xml_abas[nome] = z.read(f"xl/worksheets/sheet{i}.xml").decode("utf-8")


def celula_crua(nome_aba, coluna, codigo):
    """(atributos da tag <c>, valor <v>) direto do XML da planilha - o
    openpyxl converteria uma célula de tempo em timedelta."""
    aba = livro[nome_aba]
    cab = [x.value for x in aba[1]]
    letra = get_column_letter(cab.index(coluna) + 1)
    linha = next(r[0].row for r in aba.iter_rows(min_row=2) if r[0].value == codigo)
    m = re.search(rf'<c r="{letra}{linha}"([^>]*)>(?:<v>([^<]*)</v>)?</c>', xml_abas[nome_aba])
    assert m, (nome_aba, coluna, codigo)
    return m.group(1), m.group(2), aba[f"{letra}{linha}"]


def eh_numero(atributos):
    """Tipo da célula no XML: sem `t` ou `t="n"` = número; `t="s"`/`"str"`/
    `"inlineStr"` seria texto."""
    tipo = re.search(r' t="([^"]*)"', atributos)
    return tipo is None or tipo.group(1) == "n"


for coluna, segundos in (("Espera inicial (h:mm:ss)", 1308), ("Atendimento total (h:mm:ss)", 58524)):
    atributos, valor, cel = celula_crua("Por senha", coluna, "A")
    assert eh_numero(atributos), atributos
    assert abs(float(valor) - segundos / 86400) < 1e-12, (coluna, valor)
    assert cel.number_format == "[h]:mm:ss" and cel.value == timedelta(seconds=segundos)
atributos, valor, cel = celula_crua("Por etapa", "Duração (h:mm:ss)", "A")
assert eh_numero(atributos) and abs(float(valor) - 1308 / 86400) < 1e-12 and cel.number_format == "[h]:mm:ss"
assert "Espera inicial (min)" not in [x.value for x in livro["Por senha"][1]]
# Valor gravado = SEGUNDOS ARREDONDADOS / 86400 (decisão do 10.75): a célula mostra o
# mesmo segundo da tela, não os segundos exatos (o cenário acima só tem inteiros).
assert _duracao_excel(1307.6) == 1308 / 86400 and _duracao_excel(1307.4) == 1307 / 86400
assert _duracao_excel(20.5) == 21 / 86400 and _duracao_excel(None) is None
assert formatar_duracao(1307.6) == "21:48"  # o mesmo segundo que a célula exibe
ok("Excel: célula numérica (t=\"n\"), valor = segundos/86400, formato [h]:mm:ss, nas abas 'Por senha' e 'Por etapa'")

# ===========================================================================
# 4) CSV: minutos decimais + a mesma duração em texto, lado a lado.
# ===========================================================================
for formato, chave, colunas in (
    ("csv", "Senha", ("Espera inicial (min)", "Espera inicial (mm:ss)", "Atendimento total (min)", "Atendimento total (mm:ss)")),
    ("csv_etapas", "Senha", ("Duração (min)", "Duração (mm:ss)")),
):
    bruto = baixar("tempos-por-senha", formato).decode("utf-8-sig")
    df = pd.read_csv(io.StringIO(bruto), sep=";", decimal=",", dtype={c: str for c in colunas if "mm:ss" in c})
    for coluna in colunas:
        assert coluna in df.columns, (formato, coluna)
    if formato == "csv":
        a = df.set_index(chave).loc["A"]
        assert a["Espera inicial (min)"] == 21.8 and a["Espera inicial (mm:ss)"] == "21:48"
        assert a["Atendimento total (min)"] == 975.4 and a["Atendimento total (mm:ss)"] == "16:15:24"
        assert df.set_index(chave).loc["B"]["Atendimento total (mm:ss)"] == "-00:05"
    else:
        assert "21:48" in list(df["Duração (mm:ss)"]) and 21.8 in list(df["Duração (min)"])
ok("CSV por senha e por etapa: coluna em minutos decimais (21,80) e a mesma duração em texto (21:48), lado a lado")

# ===========================================================================
# 5) Dashboard: texto vindo do servidor, campo numérico intacto.
# ===========================================================================
with patch("app.utils.tempo._instante_utc", return_value=datetime(2026, 10, 8, 15, 0, tzinfo=UTC)):
    kpis = c.get("/api/dashboard/kpis").get_json()
    dashboard = c.get("/dashboard").get_data(as_text=True)
assert kpis["tempo_medio_espera_segundos"] == 684 and kpis["tempo_medio_espera_texto"] == "11:24", kpis
assert '"11:24" || "—"' in dashboard and "formatarTempo" not in dashboard
with app.app_context():
    from app.core.dashboard_queries import calcular_kpis
    vazio = calcular_kpis(db.session.get(Empresa, ids["empresa"]), [])
assert vazio["tempo_medio_espera_segundos"] is None and vazio["tempo_medio_espera_texto"] is None
ok("Dashboard: /api/dashboard/kpis mantém tempo_medio_espera_segundos (684) e ganha tempo_medio_espera_texto ('11:24'); sem dado = None")

print(f"\n{len(passos)} verificações OK")
