"""
Teste de regressão: relatório "Tempos por senha" (adendo 10.73).

Cenário (relógio simulado em 09/10/2026 01:00 UTC = 08/10 22:00 de Brasília;
filtro De/Até = 08/10):
  Unidade "Nova" (logs com snapshot de fila, formato do 10.62):
    S1  simples: espera 5, atendimento 7
    S2  Recepção -> Triagem -> Coleta (2 realocações): esperas 3+6+6,
        atendimentos 1+4+9, 3 etapas
    S3  nunca chamada
    S4  chamada e cancelada pelo encerramento automático (10.18)
    S5  chamada, não finalizada
    S8  emitida às 21h50 de Brasília (00:50 UTC do dia 09)
  Unidade "Histórica" (logs SEM snapshot, anteriores ao 10.62):
    S6  com realocação -> "Histórico parcial" (só espera inicial)
    S7  sem realocação -> "Concluída" (fallback exato)
Todas com nome_paciente preenchido - não pode aparecer em nenhuma saída.

Execução: python -m tests.test_regressao_tempos_por_senha (a partir da raiz
do projeto, com o virtualenv ativado).
"""
import os

os.environ.setdefault("SECRET_KEY", "test")
os.environ.setdefault("FLASK_ENV", "testing")

import html
import io
import re
from codecs import BOM_UTF8
from datetime import datetime, timedelta
from unittest.mock import patch
from zoneinfo import ZoneInfo

import openpyxl
import pandas as pd
from pypdf import PdfReader

from app import create_app
from app.extensions import db
from app.models import Empresa, Fila, Local, LogAtendimento, LogExportacaoRelatorio, Perfil, Senha, Unidade, Usuario
from app.relatorios import routes as rotas_relatorios
from app.relatorios.helpers import Filtros
from app.relatorios.queries import (
    SITUACAO_CANCELADA, SITUACAO_CONCLUIDA, SITUACAO_HISTORICO_PARCIAL, SITUACAO_NAO_FINALIZADA,
    SITUACAO_NUNCA_CHAMADA, tempo_atendimento, tempo_medio_espera, tempos_por_senha,
)
from app.utils.tempo import janela_dia_clinica_em_utc

UTC = ZoneInfo("UTC")
SP = ZoneInfo("America/Sao_Paulo")
NOME = "Maria Sigilosa"
app = create_app("testing")

passos = []


def ok(msg):
    passos.append(msg)
    print("OK:", msg)


def relogio():
    return patch("app.utils.tempo._instante_utc", return_value=datetime(2026, 10, 9, 1, 0, tzinfo=UTC))


T0 = datetime(2026, 10, 8, 12, 0)  # 09:00 em Brasília

with app.app_context():
    db.create_all()
    p_admin, p_atend = Perfil(nome="admin"), Perfil(nome="atendente")
    empresa = Empresa(nome="Clinica Tempos")
    db.session.add_all([p_admin, p_atend, empresa])
    db.session.commit()
    nova = Unidade(empresa_id=empresa.id, nome="Nova", sigla="NOV")
    hist = Unidade(empresa_id=empresa.id, nome="Historica", sigla="HIS")
    db.session.add_all([nova, hist])
    db.session.commit()

    def fila(unidade, nome):
        f = Fila(empresa_id=empresa.id, unidade_id=unidade.id, nome=nome)
        db.session.add(f)
        db.session.commit()
        return f

    f_rec, f_tri, f_col = fila(nova, "Recepção"), fila(nova, "Triagem"), fila(nova, "Coleta")
    f_hist = fila(hist, "Recepção")
    l_rec = Local(empresa_id=empresa.id, unidade_id=nova.id, nome="G1", nome_exibicao="G1", identificador_admin="[NOV] Guichê 1")
    l_col = Local(empresa_id=empresa.id, unidade_id=nova.id, nome="C1", nome_exibicao="C1", identificador_admin="[NOV] Coleta 1")
    db.session.add_all([l_rec, l_col])
    db.session.commit()

    def usuario(login, perfil, unidades):
        u = Usuario(empresa_id=empresa.id, nome=login, login=login)
        u.set_senha("senha123")
        u.perfis = [perfil]
        u.unidades = unidades
        db.session.add(u)
        db.session.commit()
        return u

    admin = usuario("admin_tempos", p_admin, [])
    op_r, op_t, op_c = (usuario(n, p_atend, [nova, hist]) for n in ("op_recepcao", "op_triagem", "op_coleta"))

    def senha(codigo, unidade, fila_atual, status, criada=T0, local=None):
        s = Senha(empresa_id=empresa.id, unidade_id=unidade.id, fila_id=fila_atual.id, codigo=codigo, sequencial=1,
                  prioridade=False, status=status, criada_em=criada, local_id=local.id if local else None,
                  nome_paciente=NOME)
        db.session.add(s)
        db.session.commit()
        return s

    def log(s, evento, minuto, usuario=None, fila_snap=None, local_snap=None, base=T0):
        db.session.add(LogAtendimento(
            senha_id=s.id, evento=evento, timestamp=base + timedelta(minutes=minuto),
            usuario_id=usuario.id if usuario else None, unidade_id=s.unidade_id,
            fila_id=fila_snap.id if fila_snap else None, local_id=local_snap.id if local_snap else None,
            setor="recepcao" if usuario is not None else None,
        ))

    s1 = senha("S1", nova, f_rec, "finalizada", local=l_rec)
    log(s1, "gerada", 0, None, f_rec); log(s1, "chamada_senha", 5, op_r, f_rec, l_rec); log(s1, "fim_atendimento", 12, op_r, f_rec, l_rec)

    s2 = senha("S2", nova, f_col, "finalizada", local=l_col)
    log(s2, "gerada", 0, None, f_rec)
    log(s2, "chamada_senha", 3, op_r, f_rec, l_rec); log(s2, "realocado", 4, op_r, f_rec, l_rec)
    log(s2, "chamada_senha", 10, op_t, f_tri); log(s2, "realocado", 14, op_t, f_tri)
    log(s2, "chamada_nome", 20, op_c, f_col, l_col); log(s2, "fim_atendimento", 29, op_c, f_col, l_col)

    s3 = senha("S3", nova, f_rec, "aguardando")
    log(s3, "gerada", 30, None, f_rec)

    s4 = senha("S4", nova, f_rec, "cancelada")
    log(s4, "gerada", 0, None, f_rec); log(s4, "chamada_senha", 8, op_r, f_rec, l_rec)
    log(s4, "encerramento_automatico_sistema", 600, None, f_rec)

    s5 = senha("S5", nova, f_rec, "chamada", local=l_rec)
    log(s5, "gerada", 0, None, f_rec); log(s5, "chamada_senha", 2, op_r, f_rec, l_rec)

    s8 = senha("S8", nova, f_rec, "finalizada", criada=datetime(2026, 10, 9, 0, 50), local=l_rec)
    base8 = datetime(2026, 10, 9, 0, 50)
    log(s8, "gerada", 0, None, f_rec, base=base8); log(s8, "chamada_senha", 5, op_r, f_rec, l_rec, base=base8)
    log(s8, "fim_atendimento", 8, op_r, f_rec, l_rec, base=base8)

    # Histórico (sem snapshot de fila em NENHUM log).
    s6 = senha("S6", hist, f_hist, "finalizada")
    log(s6, "gerada", 0); log(s6, "chamada_senha", 6, op_r); log(s6, "realocado", 7, op_r)
    log(s6, "chamada_senha", 8, op_c); log(s6, "fim_atendimento", 20, op_c)
    s7 = senha("S7", hist, f_hist, "finalizada")
    log(s7, "gerada", 0); log(s7, "chamada_senha", 4, op_r); log(s7, "fim_atendimento", 10, op_r)
    db.session.commit()

    ids = dict(empresa=empresa.id, nova=nova.id, hist=hist.id, op_t=op_t.id, op_c=op_c.id,
               f_rec=f_rec.id, f_col=f_col.id)
    INICIO, FIM = janela_dia_clinica_em_utc(datetime(2026, 10, 8, tzinfo=SP))


def filtros(unidade_id=None, operador_id=None):
    return Filtros(empresa_id=ids["empresa"], unidades_permitidas=None, unidade_id=unidade_id, fila_id=None,
                   operador_id=operador_id, data_inicio=INICIO, data_fim_exclusivo=FIM,
                   data_inicio_str="2026-10-08", data_fim_str="2026-10-08")


def por_codigo(linhas):
    return {l["codigo"]: l for l in linhas}


def minutos(segundos):
    return None if segundos is None else round(segundos / 60, 2)


# ===========================================================================
# 1) Consulta: senha simples, 2 realocações e cada caso sem dado.
# ===========================================================================
with app.app_context():
    linhas, etapas = tempos_por_senha(filtros())
L = por_codigo(linhas)
assert set(L) == {"S1", "S2", "S3", "S4", "S5", "S6", "S7", "S8"}, set(L)

def tempos(c):
    l = L[c]
    return (minutos(l["espera_inicial_segundos"]), minutos(l["espera_total_segundos"]),
            minutos(l["atendimento_total_segundos"]), l["etapas"], l["situacao"])

assert tempos("S1") == (5, 5, 7, 1, SITUACAO_CONCLUIDA), tempos("S1")
ok("Senha simples: espera inicial 5, espera total 5, atendimento 7, 1 etapa, Concluída")

assert tempos("S2") == (3, 15, 14, 3, SITUACAO_CONCLUIDA), tempos("S2")
assert L["S2"]["fila_inicial_id"] == ids["f_rec"] and L["S2"]["fila_final_id"] == ids["f_col"]
assert L["S2"]["operador_id"] == ids["op_c"]
etapas_s2 = [(e["numero"], e["tipo"], minutos(e["segundos"])) for e in etapas if e["codigo"] == "S2"]
assert etapas_s2 == [(1, "espera", 3), (2, "repasse", 1), (3, "espera", 6),
                     (4, "repasse", 4), (5, "espera", 6), (6, "conclusao", 9)], etapas_s2
ok("2 realocações: esperas 3+6+6=15, atendimentos 1+4+9=14, 3 etapas; fila inicial Recepção, final Coleta, operador de finalização")
ok("Visão por etapa da S2: 6 linhas em ordem (espera/repasse/espera/repasse/espera/conclusão)")

assert tempos("S3") == (None, None, None, None, SITUACAO_NUNCA_CHAMADA), tempos("S3")
assert tempos("S4") == (8, 8, None, None, SITUACAO_CANCELADA), tempos("S4")
assert tempos("S5") == (2, 2, None, None, SITUACAO_NAO_FINALIZADA), tempos("S5")
ok("Sem dado = vazio (None), nunca 0: nunca chamada, cancelada (encerramento automático), não finalizada")

assert tempos("S6") == (6, None, None, None, SITUACAO_HISTORICO_PARCIAL), tempos("S6")
assert L["S6"]["fila_inicial_id"] is None
assert tempos("S7") == (4, 4, 6, 1, SITUACAO_CONCLUIDA), tempos("S7")
ok("Histórico com realocação = 'Histórico parcial' (só espera inicial); histórico sem realocação = 'Concluída'")

# ===========================================================================
# 2) Coerência com os agregados (só dado no formato novo - unidade Nova).
# ===========================================================================
with app.app_context():
    linhas_n, etapas_n = tempos_por_senha(filtros(ids["nova"]))
    agg_espera = tempo_medio_espera(filtros(ids["nova"]))
    agg_atend = tempo_atendimento(filtros(ids["nova"]))

soma_espera = sum(l["espera_total_segundos"] for l in linhas_n if l["espera_total_segundos"] is not None)
soma_agg_espera = sum(r["quantidade"] * r["tempo_medio_segundos"] for r in agg_espera)
qtd_agg_espera = sum(r["quantidade"] for r in agg_espera)
esperas_etapa = [e["segundos"] for e in etapas_n if e["tipo"] == "espera"]
assert abs(soma_espera - soma_agg_espera) < 0.1 * qtd_agg_espera, (soma_espera, soma_agg_espera)
assert len(esperas_etapa) == qtd_agg_espera == 7, (len(esperas_etapa), qtd_agg_espera)
assert abs(sum(esperas_etapa) / len(esperas_etapa) - soma_agg_espera / qtd_agg_espera) < 0.1
ok("Espera: soma por senha = soma do agregado, 7 etapas = 7 do agregado, média por etapa = média do agregado")

soma_atend = sum(l["atendimento_total_segundos"] for l in linhas_n if l["atendimento_total_segundos"] is not None)
soma_agg_atend = sum(r["quantidade"] * r["tempo_medio_segundos"] for r in agg_atend)
qtd_agg_atend = sum(r["quantidade"] for r in agg_atend)
atend_etapa = [e["segundos"] for e in etapas_n if e["tipo"] != "espera"]
assert abs(soma_atend - soma_agg_atend) < 0.1 * qtd_agg_atend, (soma_atend, soma_agg_atend)
assert len(atend_etapa) == qtd_agg_atend == 5
assert abs(sum(atend_etapa) / len(atend_etapa) - soma_agg_atend / qtd_agg_atend) < 0.1
ok("Atendimento: soma por senha = soma do agregado, 5 etapas = 5 do agregado, média por etapa = média do agregado")

# ===========================================================================
# 3) Tela (relógio às 22h), resumo, tooltip, sem nome do paciente.
# ===========================================================================
c = app.test_client()
assert c.post("/login", data={"login": "admin_tempos", "senha": "senha123"}, follow_redirects=True).status_code == 200
URL = "/relatorios/tempos-por-senha?data_inicio=2026-10-08&data_fim=2026-10-08"


def get(url):
    with relogio():
        return c.get(url)


tela = html.unescape(get(URL).get_data(as_text=True))
assert "<td>08/10/2026 21:50</td>" in tela  # S8 no fuso da clínica e dentro do dia 08/10
assert f'title="{rotas_relatorios.DICAS_SITUACAO[SITUACAO_HISTORICO_PARCIAL]}">Histórico parcial</span>' in tela
assert "7 de 8 senhas" in tela  # espera inicial: só S3 fica fora
assert "Média e mediana POR SENHA" in tela and "fazem a média POR ETAPA" in tela
assert "<th>Etapas de atendimento</th>" in tela and '"Operador (chamada/atendimento)"' in tela
assert NOME not in tela and "Sigilosa" not in tela
ok("Tela: emissão 21:50 no dia 08/10 (relógio às 22h), tooltip do Histórico parcial, '7 de 8 senhas', nota da média")
ok("Tela sem nome do paciente")

with relogio(), patch.object(rotas_relatorios, "LINHAS_POR_PAGINA", 3):
    p1 = html.unescape(c.get(URL).get_data(as_text=True))
    p3 = html.unescape(c.get(URL + "&pagina=3").get_data(as_text=True))
assert "Página 1 de 3 (8 senhas)" in p1 and "pagina=2" in p1
assert p1.count("<td>S") == 3 and p3.count("<td>S") == 2
assert "7 de 8 senhas" in p3  # resumo sempre sobre o filtro inteiro, não só a página
ok("Paginação (LINHAS_POR_PAGINA): 3 por página, resumo calculado sobre todas as 8 senhas")

# ===========================================================================
# 4) Exportação: números (não texto), CSV pt-BR, Excel em duas abas, sem nome.
# ===========================================================================
def baixar(formato, extra=""):
    resp = get(f"{URL}&formato={formato}{extra}")
    assert resp.status_code == 200, (formato, resp.status_code)
    resp.direct_passthrough = False
    return resp.get_data()

bruto_bytes = baixar("csv")
assert bruto_bytes.startswith(BOM_UTF8)  # UTF-8 COM BOM: o Excel pt-BR não quebra os acentos
bruto = bruto_bytes.decode("utf-8-sig")
assert "Recepção" in bruto
assert NOME not in bruto and ";" in bruto.splitlines()[0]
df = pd.read_csv(io.StringIO(bruto), sep=";", decimal=",")
assert df["Espera inicial (min)"].dtype == float and df["Atendimento total (min)"].dtype == float
linha_s2 = bruto.splitlines()[[l.split(";")[0] for l in bruto.splitlines()].index("S2")]
# Minutos decimais + o mesmo tempo em texto (adendo 10.75); etapas como inteiro ("3", não "3,00").
assert ";3,00;03:00;15,00;15:00;14,00;14:00;3;" in linha_s2, linha_s2
assert df.set_index("Senha").loc["S3"].isna()["Espera inicial (min)"]
assert "Etapas de atendimento" in df.columns
ok("CSV por senha: UTF-8 com BOM, padrão pt-BR (';' e ','), minutos com 2 casas como número; sem dado = vazio; sem nome")

bytes_etapas = baixar("csv_etapas", f"&unidade_id={ids['nova']}")
assert bytes_etapas.startswith(BOM_UTF8)
bruto_etapas = bytes_etapas.decode("utf-8-sig")
df_e = pd.read_csv(io.StringIO(bruto_etapas), sep=";", decimal=",")
assert len(df_e) == 12 and df_e["Duração (min)"].dtype == float and NOME not in bruto_etapas
assert "Operador (chamada/atendimento)" in df_e.columns and "Operador" not in df_e.columns
ok("CSV por etapa: 12 etapas da unidade Nova (7 esperas + 5 atendimentos), minutos numéricos, sem nome")

livro = openpyxl.load_workbook(io.BytesIO(baixar("xlsx")))
assert livro.sheetnames == ["Por senha", "Por etapa"], livro.sheetnames
assert "Etapas de atendimento" in [c.value for c in livro["Por senha"][1]]
assert "Operador (chamada/atendimento)" in [c.value for c in livro["Por etapa"][1]]
aba = livro["Por senha"]
cab = [c.value for c in aba[1]]
col = {nome: i for i, nome in enumerate(cab)}
linhas_xlsx = {r[0].value: r for r in aba.iter_rows(min_row=2)}
# Duração nativa (adendo 10.75): o openpyxl lê uma célula numérica com formato de
# tempo como timedelta; o valor numérico cru (segundos/86400) é conferido no XML
# em tests/test_regressao_duracao.py.
celula = linhas_xlsx["S2"][col["Espera total (h:mm:ss)"]]
assert celula.value == timedelta(minutes=15) and celula.number_format == "[h]:mm:ss", (celula.value, celula.number_format)
assert linhas_xlsx["S1"][col["Espera inicial (h:mm:ss)"]].value == timedelta(minutes=5)
emissao = linhas_xlsx["S8"][col["Emissão"]]
assert emissao.value == datetime(2026, 10, 8, 21, 50) and emissao.number_format == "dd/mm/yyyy hh:mm"
assert linhas_xlsx["S3"][col["Espera inicial (h:mm:ss)"]].value is None
for nome_aba in livro.sheetnames:
    for linha in livro[nome_aba].iter_rows(values_only=True):
        assert NOME not in [str(v) for v in linha], nome_aba
ok("Excel: abas 'Por senha'/'Por etapa', tempos como duração [h]:mm:ss, emissão como data/hora da clínica, sem nome")

# Filtro de operador: (i) por senha mantém a senha inteira, (ii) por etapa só as do operador.
with relogio():
    tela_op = html.unescape(c.get(f"{URL}&operador_id={ids['op_t']}").get_data(as_text=True))
assert tela_op.count("<td>S") == 1 and "<td>S2</td>" in tela_op and "<td>14:00</td>" in tela_op
df_op = pd.read_csv(io.StringIO(baixar("csv_etapas", f"&operador_id={ids['op_t']}").decode("utf-8-sig")), sep=";", decimal=",")
assert list(zip(df_op["Etapa"], df_op["Tipo"], df_op["Duração (min)"])) == [(3, "Espera", 6.0), (4, "Atendimento - Repasse", 4.0)], df_op
ok("Operador: por senha mantém S2 com o total inteiro (14 min); por etapa só as 2 etapas do operador (numeração preservada)")

# ===========================================================================
# 5) PDF (Chromium real): dados, limite de linhas, sem nome; log de auditoria.
# ===========================================================================
def texto_pdf(conteudo):
    # Espaços normalizados: colunas estreitas quebram "Histórico parcial" em duas linhas na extração.
    return " ".join("".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(conteudo)).pages).split())

pdf = texto_pdf(baixar("pdf"))
assert "S2" in pdf and "Histórico parcial" in pdf and NOME not in pdf
assert "Etapas de atendimento" in pdf
with patch.object(rotas_relatorios, "LIMITE_LINHAS_PDF", 2):
    pdf_resumo = texto_pdf(baixar("pdf"))
assert "acima do limite de 2 linhas do PDF" in pdf_resumo and "Média:" in pdf_resumo
assert not re.search(r"\bS2\b", pdf_resumo)
ok("PDF: lista completa abaixo do limite; acima de LIMITE_LINHAS_PDF só o resumo + aviso de CSV/Excel; sem nome")

with app.app_context():
    registros = {(r.tipo_relatorio, r.formato) for r in LogExportacaoRelatorio.query.all()}
assert {("tempos_por_senha", "csv"), ("tempos_por_senha", "xlsx"), ("tempos_por_senha", "pdf"),
        ("tempos_por_senha_etapas", "csv")} <= registros, registros
ok("Log de auditoria (10.67) para CSV, Excel, PDF e CSV por etapa")

# ===========================================================================
# 6) RBAC e menu.
# ===========================================================================
c_op = app.test_client()
c_op.post("/login", data={"login": "op_recepcao", "senha": "senha123"})
assert c_op.get(URL).status_code == 403
assert "Tempos por senha" in get("/relatorios/").get_data(as_text=True)
ok("RBAC igual aos relatórios (403 para atendente) e item 'Tempos por senha' no menu")

print(f"\n{len(passos)} verificações OK")
